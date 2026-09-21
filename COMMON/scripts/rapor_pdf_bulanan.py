"""
Rapor nilai bulanan (DKN / B5 UTS1) 2026.xlsm -> 1 PDF gabungan per kelas.
Pakai:  py rapor_pdf_bulanan.py "<path xlsm>" [kode kelas]   (tanpa kelas = semua)
Contoh: py rapor_pdf_bulanan.py "C:\RaporServer\...\B5 UTS1\2026.xlsm" VB SD
Mapping sheet: kelas 1/2 -> A, 3/4 -> B, 5/6 -> C.
Output: <KELAS>.pdf di folder yang sama dengan xlsm (temp per murid dibersihkan).
"""
import os, re, sys, tempfile, shutil, time

import win32com.client
import pythoncom
import openpyxl
import pymupdf

def com_call(fn, *args, tries=60):
    for _ in range(tries):
        try:
            return fn(*args)
        except pythoncom.com_error as e:
            if e.hresult in (-2147418111, -2147417846):  # rejected/busy
                time.sleep(0.5)
                continue
            raise
    raise SystemExit('Excel tetap sibuk (call rejected berulang)')

def detect_sheet(kelas):
    # romawi -> sheet (kelas 1/2 -> A, 3/4 -> B, 5/6 -> C)
    rules = [('III', 'B'), ('IV', 'B'), ('VI', 'C'),
             ('II', 'A'), ('V', 'C'), ('I', 'A')]
    for rom, s in rules:
        if kelas.startswith(rom):
            return s
    raise SystemExit(f'Tidak bisa deteksi sheet untuk kelas {kelas!r}')

def kelas_murid(ws_data, kelas):
    """Daftar nomor yang benar-benar ada utk kelas (bisa ada gap).
    Kunci Data pola: '<KELAS> SD<n>' contoh 'IA SD1' (spasi bisa 1-2)."""
    pat = re.compile(r'^\s*' + re.escape(kelas) + r'\s*(\d+)$', re.IGNORECASE)
    nums = []
    for r in ws_data.iter_rows(min_row=3, max_col=1):
        v = r[0].value
        if isinstance(v, str):
            m = pat.match(v.strip())
            if m:
                nums.append(int(m.group(1)))
    return sorted(set(nums))

def semua_kelas(ws_data):
    pat = re.compile(r'^\s*([IVXL]+\s*[AB])\s+SD\s*\d+$', re.IGNORECASE)
    kelas = {}
    for r in ws_data.iter_rows(min_row=3, max_col=1):
        v = r[0].value
        if isinstance(v, str):
            m = pat.match(v.strip())
            if m:
                k = m.group(1).upper().replace(' ', '') + ' SD'
                kelas[k] = kelas.get(k, 0) + 1
    return sorted(kelas)

def main():
    if len(sys.argv) < 2:
        raise SystemExit('Pakai: py rapor_pdf_bulanan.py "<xlsm>" [kelas]')
    xlsm = sys.argv[1]
    if not os.path.isfile(xlsm):
        raise SystemExit(f'File tidak ditemukan: {xlsm}')
    dir_ = os.path.dirname(xlsm)

    # in-sheet macros tidak relevan: kita buka langsung (macro aktif tidak wajib utk export)
    wb_info = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    ws_data = wb_info['Data']

    if len(sys.argv) >= 3:
        targets = [sys.argv[2].strip()]
    else:
        targets = semua_kelas(ws_data)
    print(f'Kelas yang akan diproses: {targets}')

    excel = None; wb = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        wb = com_call(excel.Workbooks.Open, xlsm, False, True)
        pythoncom.CoInitialize()

        for kelas in targets:
            sheet = detect_sheet(kelas)
            nums = kelas_murid(ws_data, kelas)
            if not nums:
                print(f'  {kelas}: tidak ada data, dilewati')
                continue
            if len(nums) != nums[-1] - nums[0] + 1:
                missing = [n for n in range(nums[0], nums[-1] + 1) if n not in set(nums)]
                print(f'  {kelas}: catatan nomor kosong di Data: {missing} (dilewati)')
            print(f'  {kelas} -> sheet {sheet}, murid {len(nums)} (no {nums[0]}-{nums[-1]})')

            ws = com_call(wb.Worksheets, sheet)
            com_call(ws.Activate)
            ws.Range('F2').Value = kelas
            ws.Range('H7').Value = nums[0]
            ws.Range('H8').Value = nums[-1]

            tmp = tempfile.mkdtemp(prefix='rapor_')
            files = []
            try:
                for i in sorted(nums, reverse=True):
                    ws.Range('H1').Value = i
                    try:
                        com_call(ws.Calculate)
                    except Exception:
                        pass
                    no = ws.Range('D7').Value
                    if isinstance(no, float) and no.is_integer():
                        no = int(no)
                    nama = ws.Range('D9').Value
                    if nama is None or str(nama).strip() == '':
                        print(f'    skip no {i}: nama kosong (bukan murid)')
                        continue
                    nama = str(nama)
                    safe = re.sub(r'[\\/:*?"<>|]', '_', f'{no}-{nama}-{i}')
                    out = os.path.join(tmp, f'{i:03d}.pdf')
                    com_call(ws.ExportAsFixedFormat, 0, out)
                    files.append((i, out, safe))
                    print(f'    {safe}')

                merged = pymupdf.open()
                for i, out, safe in sorted(files, key=lambda x: x[0]):
                    src = pymupdf.open(out)
                    merged.insert_pdf(src)
                    src.close()
                dest = os.path.join(dir_, f'{kelas}.pdf')
                merged.save(dest, garbage=4, deflate=True)
                pages = merged.page_count
                merged.close()
                print(f'  OK: {dest}  ({pages} halaman)')
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

if __name__ == '__main__':
    main()