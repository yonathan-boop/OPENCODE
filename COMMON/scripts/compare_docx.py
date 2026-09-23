#!/usr/bin/env python3
"""
Standard Document Comparison & Redline Tool for Word (.docx).
Uses Word / WPS Office COM engine to generate native Track Changes comparison documents
(strikethrough for deletions, colored underline for insertions), and audits leftover placeholders.
"""

import sys
import os
import argparse
import re
import docx
import win32com.client

def audit_placeholders(docx_path):
    """Scans docx for lingering placeholders, templates, and incomplete fields."""
    patterns = [
        (r'\.{3,}', 'Titik-titik placeholder (...)'),
        (r'…+', 'Karakter ellipsis (…)'),
        (r'\bUPT SD\b', 'Sisa template UPT SD'),
        (r'\bSDN\s*/\s*SDS\b', 'Sisa template SDN / SDS'),
        (r'\(.*uraikan.*\)', 'Instruksi (uraikan...)'),
        (r'\(.*tambahkan.*\)', 'Instruksi (tambahkan...)'),
        (r'\(.*contoh.*\)', 'Instruksi (contoh...)'),
        (r'\[.*\]', 'Tanda kurung siku placeholder [...]')
    ]
    
    issues = []
    try:
        doc = docx.Document(docx_path)
        for idx, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            if not txt:
                continue
            for pat, desc in patterns:
                if re.search(pat, txt, re.IGNORECASE):
                    issues.append(f"[Paragraf {idx+1}] {desc}: \"{txt[:100]}...\"")
                    break

        for t_idx, tbl in enumerate(doc.tables):
            for r_idx, row in enumerate(tbl.rows):
                for c_idx, cell in enumerate(row.cells):
                    ctxt = cell.text.strip()
                    if not ctxt:
                        continue
                    for pat, desc in patterns:
                        if re.search(pat, ctxt, re.IGNORECASE):
                            issues.append(f"[Tabel {t_idx+1} B{r_idx+1}K{c_idx+1}] {desc}: \"{ctxt[:100]}...\"")
                            break
    except Exception as e:
        issues.append(f"Gagal memindai placeholder: {e}")
        
    return issues

def compare_documents(original_path, revised_path, output_path):
    """Compares original_path and revised_path using Word/WPS COM."""
    orig_abs = os.path.abspath(original_path)
    rev_abs = os.path.abspath(rev_path)
    out_abs = os.path.abspath(output_path)

    if not os.path.exists(orig_abs):
        raise FileNotFoundError(f"File asli tidak ditemukan: {orig_abs}")
    if not os.path.exists(rev_abs):
        raise FileNotFoundError(f"File revisi tidak ditemukan: {rev_abs}")

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    try:
        doc_orig = word.Documents.Open(orig_abs)
        doc_rev = word.Documents.Open(rev_abs)

        # wdCompareDestinationNew = 2, wdGranularityWordLevel = 1
        doc_diff = word.CompareDocuments(
            OriginalDocument=doc_orig,
            RevisedDocument=doc_rev,
            Destination=2,
            Granularity=1,
            CompareFormatting=False,
            AddToMru=False
        )

        doc_diff.SaveAs(out_abs)
        revisions_count = doc_diff.Revisions.Count

        doc_diff.Close(False)
        doc_orig.Close(False)
        doc_rev.Close(False)
        return revisions_count
    finally:
        try:
            word.Quit()
        except Exception:
            pass

def main():
    parser = argparse.ArgumentParser(description="Word DOCX Compare & Redline Utility")
    parser.add_argument("original", help="Path ke dokumen template / asli")
    parser.add_argument("revised", help="Path ke dokumen hasil edit")
    parser.add_argument("-o", "--output", required=True, help="Path file output hasil perbandingan")
    parser.add_argument("--no-audit", action="store_true", help="Lewati pemeriksaan placeholder")

    args = parser.parse_args()

    print(f"[*] Membandingkan:")
    print(f"    Asli   : {args.original}")
    print(f"    Revisi : {args.revised}")
    print(f"    Output : {args.output}")

    try:
        rev_count = compare_documents(args.original, args.revised, args.output)
        print(f"[+] Berhasil! Dihasilkan file perbandingan dengan {rev_count} revisi tercatat.")
    except Exception as e:
        print(f"[-] Terjadi kesalahan saat membandingkan dokumen: {e}")
        sys.exit(1)

    if not args.no_audit:
        print("\n[*] Memindai sisa placeholder / template yang belum terisi di dokumen revisi...")
        issues = audit_placeholders(args.revised)
        if issues:
            print(f"[!] Ditemukan {len(issues)} potensi placeholder yang belum selesai:")
            for issue in issues[:15]:
                print(f"    - {issue}")
            if len(issues) > 15:
                print(f"    ... dan {len(issues)-15} lainnya.")
        else:
            print("[+] Bersih! Tidak ditemukan placeholder template yang tertinggal.")

if __name__ == "__main__":
    main()
